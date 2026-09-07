import os
from io import BytesIO
import pandas as pd
import pdfplumber
import pptx
from docx import Document
from PyPDF2 import PdfReader
import chromadb
from chromadb.config import Settings
from langchain_text_splitters import RecursiveCharacterTextSplitter

# Initialize ChromaDB client (local persistent)
# In production, you might want this to point to a Chroma server
chroma_client = chromadb.PersistentClient(path="./chroma_db")

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200

def _extract_pdf(file_content: bytes) -> str:
    text = ""
    try:
        with pdfplumber.open(BytesIO(file_content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text += page_text + "\n"
                tables = page.extract_tables()
                for table in tables or []:
                    rows = ["\t".join(cell if cell is not None else "" for cell in row) for row in table]
                    text += "\n".join(rows) + "\n"
    except Exception as e:
        print(f"pdfplumber failed, falling back to PyPDF2: {e}")
        pdf_reader = PdfReader(BytesIO(file_content))
        for page in pdf_reader.pages:
            text += page.extract_text() or ""
    return text

def _extract_docx(file_content: bytes) -> str:
    doc = Document(BytesIO(file_content))
    return "".join(para.text + "\n" for para in doc.paragraphs)

def _extract_pptx(file_content: bytes) -> str:
    ppt = pptx.Presentation(BytesIO(file_content))
    text = ""
    for slide in ppt.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def _extract_excel(file_content: bytes) -> str:
    df = pd.read_excel(BytesIO(file_content))
    return df.to_string()

def extract_text_from_file(file_content: bytes, filename: str) -> str:
    """Extracts text based on file type."""
    file_ext = os.path.splitext(filename)[1].lower()
    if file_ext == ".pdf":
        return _extract_pdf(file_content)
    if file_ext == ".docx":
        return _extract_docx(file_content)
    if file_ext == ".txt":
        return file_content.decode("utf-8")
    if file_ext in [".xlsx", ".xls"]:
        return _extract_excel(file_content)
    if file_ext in [".pptx", ".ppt"]:
        return _extract_pptx(file_content)
    return ""

def create_vector_store(document_id: str, text: str):
    """Create chunks and store them in ChromaDB."""
    text_splitter = RecursiveCharacterTextSplitter(
        separators=["\n"],
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        length_function=len
    )
    chunks = text_splitter.split_text(text)
    
    collection = chroma_client.get_or_create_collection(name=document_id)
    
    ids = [f"{document_id}_{i}" for i in range(len(chunks))]
    metadatas = [{"source": document_id} for _ in range(len(chunks))]
    
    collection.add(
        documents=chunks,
        metadatas=metadatas,
        ids=ids
    )
    return collection.count()

def query_vector_store(document_id: str, query: str, n_results: int = 6):
    """Retrieve relevant chunks from ChromaDB."""
    try:
        collection = chroma_client.get_collection(name=document_id)
        results = collection.query(
            query_texts=[query],
            n_results=n_results
        )
        if results and results['documents']:
            return " ".join(results['documents'][0])
        return ""
    except ValueError:
        return "" # Collection does not exist
